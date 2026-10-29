"""
ECOUNT ERP 货柜费用分摊与全套导入表生成系统 - 图形用户界面 (GUI) v2.1
包含：
1. 智能批量解析与生成
2. 单柜录入与明细核算
3. 📐 表头列映射配置与位置调整（支持动态列位置与自定义调整）
4. 品目编码主数据管理（自动替换产品名空格为 '-' 作为编码）
5. 往来单位编码配置
6. 实时运行日志与状态反馈
"""
import sys
import os
import re
import json
import threading
import datetime
import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
import process_container as core

try:
    import ctypes
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    pass


class ContainerProcessingApp:
    
    def __init__(self, root):
        self.root = root
        self.root.title('ECOUNT ERP 货柜费用分摊与导入表生成系统 v2.1')
        self.root.geometry('1160x780')
        self.root.minsize(1000, 660)
        self.center_window(1160, 780)
        self.containers = []
        self.master_items = list(core.DEFAULT_EXISTING_MASTER)
        self.counterparty_map = core.load_counterparty_map()
        self.header_mapping = core.load_header_mapping()
        self.template_config = core.load_template_config()
        self.var_export_individual = tk.BooleanVar(value = self.template_config.get('export_individual_files', False))
        self.var_export_bundle = tk.BooleanVar(value = self.template_config.get('export_bundle_workbook', True))
        self.raw_sum_headers = []
        self.raw_det_headers = []
        self.last_generated_files = []
        self.setup_styles()
        self.create_widgets()
        self.auto_detect_default_file()

    
    def center_window(self, width, height):
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        x = max(0, (screen_width - width) // 2)
        y = max(0, (screen_height - height) // 2 - 30)
        self.root.geometry(f'''{width}x{height}+{x}+{y}''')

    
    def setup_styles(self):
        self.style = ttk.Style()
        
        try:
            self.style.theme_use('clam')
        except Exception:
            pass

        font_family = '微软雅黑'
        self.root.option_add('*Font', f'''{font_family} 9''')
        self.c_primary = '#1F4E78'
        self.c_primary_hover = '#2B579A'
        self.c_success = '#107C41'
        self.c_bg = '#F4F6F9'
        self.c_card = '#FFFFFF'
        self.c_border = '#D0D7DE'
        self.root.configure(bg = self.c_bg)
        self.style.configure('TNotebook', background = self.c_bg, borderwidth = 0)
        self.style.configure('TNotebook.Tab', font = (font_family, 9, 'bold'), padding = [
            14,
            6], background = '#E1E4E8')
        self.style.map('TNotebook.Tab', background = [
            ('selected', self.c_primary),
            ('active', '#D0D7DE')], foreground = [
            ('selected', '#FFFFFF'),
            ('active', '#000000')])
        self.style.configure('Card.TFrame', background = self.c_card, relief = 'solid', borderwidth = 1)
        self.style.configure('Main.TFrame', background = self.c_bg)
        self.style.configure('Primary.TButton', font = (font_family, 9, 'bold'), background = self.c_primary, foreground = '#FFFFFF', borderwidth = 1, focuscolor = 'none')
        self.style.map('Primary.TButton', background = [
            ('active', self.c_primary_hover),
            ('disabled', '#CCCCCC')])
        self.style.configure('Success.TButton', font = (font_family, 9, 'bold'), background = self.c_success, foreground = '#FFFFFF', borderwidth = 1, focuscolor = 'none')
        self.style.map('Success.TButton', background = [
            ('active', '#0B5A2F'),
            ('disabled', '#CCCCCC')])
        self.style.configure('Treeview.Heading', font = (font_family, 9, 'bold'), background = '#EAECEF', foreground = '#24292F')
        self.style.configure('Treeview', font = (font_family, 9), rowheight = 24, borderwidth = 1)

    
    def create_widgets(self):
        top_frame = tk.Frame(self.root, bg = self.c_primary, height = 54)
        top_frame.pack(fill = 'x', side = 'top')
        top_frame.pack_propagate(False)
        title_lbl = tk.Label(top_frame, text = 'ECOUNT ERP 货柜费用分摊与导入表生成系统', font = ('微软雅黑', 13, 'bold'), bg = self.c_primary, fg = '#FFFFFF')
        title_lbl.pack(side = 'left', padx = 16, pady = 10)
        sub_lbl = tk.Label(top_frame, text = '标准5合1导入工作簿 | 凭证·入库·品目·三费分摊 | 动态表头映射', font = ('微软雅黑', 9), bg = self.c_primary, fg = '#D9E1F2')
        sub_lbl.pack(side = 'left', padx = 5, pady = 14)
        ver_lbl = tk.Label(top_frame, text = 'v2.1 增强版', font = ('微软雅黑', 9, 'italic'), bg = self.c_primary, fg = '#B0C4DE')
        ver_lbl.pack(side = 'right', padx = 16, pady = 12)
        main_container = ttk.Frame(self.root, style = 'Main.TFrame')
        main_container.pack(fill = 'both', expand = True, padx = 10, pady = 8)
        self.notebook = ttk.Notebook(main_container)
        self.notebook.pack(fill = 'both', expand = True)
        self.tab_batch = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_batch, text = ' 📑 智能批量解析与生成 ')
        self.build_tab_batch()
        self.tab_single = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_single, text = ' 📝 单柜录入与明细核算 ')
        self.build_tab_single()
        self.tab_mapping = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_mapping, text = ' 📐 表头映射与位置调整 ')
        self.build_tab_mapping()
        self.tab_template = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_template, text = ' 📋 导入模板与批注管理 ')
        self.build_tab_template()
        self.tab_master = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_master, text = ' 📦 品目编码主数据 ')
        self.build_tab_master()
        self.tab_counterparty = ttk.Frame(self.notebook, style = 'Main.TFrame')
        self.notebook.add(self.tab_counterparty, text = ' 🏢 往来单位编码配置 ')
        self.build_tab_counterparty()
        self.build_log_panel(main_container)

    
    def build_tab_batch(self):
        card_top = tk.Frame(self.tab_batch, bg = self.c_card, bd = 1, relief = 'solid')
        card_top.pack(fill = 'x', padx = 6, pady = 4)
        f_row1 = tk.Frame(card_top, bg = self.c_card)
        f_row1.pack(fill = 'x', padx = 10, pady = (8, 4))
        tk.Label(f_row1, text = '发货汇总表 Excel:', font = ('微软雅黑', 9, 'bold'), bg = self.c_card).pack(side = 'left')
        self.var_source_file = tk.StringVar()
        self.entry_source = ttk.Entry(f_row1, textvariable = self.var_source_file)
        self.entry_source.pack(side = 'left', fill = 'x', expand = True, padx = 8)
        btn_browse = ttk.Button(f_row1, text = '浏览文件...', command = self.on_browse_source_file)
        btn_browse.pack(side = 'left', padx = 3)
        btn_parse = ttk.Button(f_row1, text = '🔍 智能解析货柜', style = 'Primary.TButton', command = self.on_parse_source_file)
        btn_parse.pack(side = 'left', padx = 3)
        f_row2 = tk.Frame(card_top, bg = self.c_card)
        f_row2.pack(fill = 'x', padx = 10, pady = (4, 8))
        tk.Label(f_row2, text = '报表输出目录:    ', font = ('微软雅黑', 9, 'bold'), bg = self.c_card).pack(side = 'left')
        self.var_output_dir = tk.StringVar(value = os.path.join(core.get_base_dir(), '报表输出'))
        self.entry_output = ttk.Entry(f_row2, textvariable = self.var_output_dir)
        self.entry_output.pack(side = 'left', fill = 'x', expand = True, padx = 8)
        btn_browse_out = ttk.Button(f_row2, text = '选择目录...', command = self.on_browse_output_dir)
        btn_browse_out.pack(side = 'left', padx = 3)
        btn_open_out = ttk.Button(f_row2, text = '📂 打开输出目录', command = self.on_open_output_dir)
        btn_open_out.pack(side = 'left', padx = 3)
        card_table = tk.Frame(self.tab_batch, bg = self.c_card, bd = 1, relief = 'solid')
        card_table.pack(fill = 'both', expand = True, padx = 6, pady = 4)
        t_action_bar = tk.Frame(card_table, bg = '#F0F4F8', height = 32)
        t_action_bar.pack(fill = 'x')
        t_action_bar.pack_propagate(False)
        tk.Label(t_action_bar, text = '📋 已识别货柜清单', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        btn_sel_all = ttk.Button(t_action_bar, text = '全选', width = 6, command = self.on_select_all_containers)
        btn_sel_all.pack(side = 'left', padx = 3)
        btn_desel = ttk.Button(t_action_bar, text = '反选', width = 6, command = self.on_invert_select_containers)
        btn_desel.pack(side = 'left', padx = 3)
        btn_load_to_single = ttk.Button(t_action_bar, text = '➡️ 载入选中柜到单柜核算', command = self.on_load_selected_to_single)
        btn_load_to_single.pack(side = 'left', padx = 8)
        btn_check_mapping = ttk.Button(t_action_bar, text = '📐 调整表头列位置', command = (lambda : self.notebook.select(self.tab_mapping)))
        btn_check_mapping.pack(side = 'left', padx = 4)
        self.lbl_container_count = tk.Label(t_action_bar, text = '共 0 个货柜', font = ('微软雅黑', 9), bg = '#F0F4F8', fg = '#555555')
        self.lbl_container_count.pack(side = 'right', padx = 12)
        cols = ('sel', 'seq', 'company', 'contract', 'container', 'date', 'rate', 'ocean', 'truck', 'insurance', 'agency', 'items_count', 'cbm', 'status')
        self.tree_containers = ttk.Treeview(card_table, columns = cols, show = 'headings', selectmode = 'extended')
        self.tree_containers.heading('sel', text = '勾选')
        self.tree_containers.heading('seq', text = '序号')
        self.tree_containers.heading('company', text = '外贸公司')
        self.tree_containers.heading('contract', text = '合同号')
        self.tree_containers.heading('container', text = '集装箱号')
        self.tree_containers.heading('date', text = '出货日期')
        self.tree_containers.heading('rate', text = '汇率')
        self.tree_containers.heading('ocean', text = '海运费($)')
        self.tree_containers.heading('truck', text = '托箱费(¥)')
        self.tree_containers.heading('insurance', text = '保费($)')
        self.tree_containers.heading('agency', text = '代理/点费(¥)')
        self.tree_containers.heading('items_count', text = '品目数')
        self.tree_containers.heading('cbm', text = '总体积(CBM)')
        self.tree_containers.heading('status', text = '状态')
        self.tree_containers.column('sel', width = 45, anchor = 'center')
        self.tree_containers.column('seq', width = 45, anchor = 'center')
        self.tree_containers.column('company', width = 80, anchor = 'center')
        self.tree_containers.column('contract', width = 95, anchor = 'center')
        self.tree_containers.column('container', width = 160, anchor = 'w')
        self.tree_containers.column('date', width = 85, anchor = 'center')
        self.tree_containers.column('rate', width = 65, anchor = 'e')
        self.tree_containers.column('ocean', width = 85, anchor = 'e')
        self.tree_containers.column('truck', width = 85, anchor = 'e')
        self.tree_containers.column('insurance', width = 75, anchor = 'e')
        self.tree_containers.column('agency', width = 90, anchor = 'e')
        self.tree_containers.column('items_count', width = 60, anchor = 'center')
        self.tree_containers.column('cbm', width = 90, anchor = 'e')
        self.tree_containers.column('status', width = 75, anchor = 'center')
        sb_v = ttk.Scrollbar(card_table, orient = 'vertical', command = self.tree_containers.yview)
        sb_h = ttk.Scrollbar(card_table, orient = 'horizontal', command = self.tree_containers.xview)
        self.tree_containers.configure(yscrollcommand = sb_v.set, xscrollcommand = sb_h.set)
        self.tree_containers.pack(fill = 'both', expand = True, side = 'top')
        sb_v.pack(side = 'right', fill = 'y')
        sb_h.pack(side = 'bottom', fill = 'x')
        self.tree_containers.bind('<Double-1>', self.on_double_click_container_row)
        card_bottom = tk.Frame(self.tab_batch, bg = self.c_card, bd = 1, relief = 'solid')
        card_bottom.pack(fill = 'x', padx = 6, pady = (4, 6))
        b_opt_bar = tk.Frame(card_bottom, bg = '#F0F4F8', height = 32)
        b_opt_bar.pack(fill = 'x')
        b_opt_bar.pack_propagate(False)
        ttk.Checkbutton(b_opt_bar, text = '同时生成并导出 ECOUNT 独立上传单表 (采购入库 / 新增品目 / 一般凭证单表)', variable = self.var_export_individual).pack(side = 'left', padx = 10)
        btn_view_tpl_specs = ttk.Button(b_opt_bar, text = '📋 配置模板与查看批注规范', command = (lambda : self.notebook.select(self.tab_template)))
        btn_view_tpl_specs.pack(side = 'right', padx = 10)
        b_frame = tk.Frame(card_bottom, bg = self.c_card)
        b_frame.pack(fill = 'x', padx = 10, pady = 8)
        self.btn_batch_generate = ttk.Button(b_frame, text = '🚀 一键生成所选货柜全套ECOUNT上传表 (5合1)', style = 'Primary.TButton', command = self.on_batch_generate)
        self.btn_batch_generate.pack(side = 'left', padx = 4, ipady = 3)
        self.progress_bar = ttk.Progressbar(b_frame, orient = 'horizontal', mode = 'determinate', length = 260)
        self.progress_bar.pack(side = 'left', padx = 10, fill = 'x', expand = True)
        self.lbl_progress = tk.Label(b_frame, text = '准备就绪', font = ('微软雅黑', 9), bg = self.c_card, fg = '#555555')
        self.lbl_progress.pack(side = 'right', padx = 8)

    
    def build_tab_single(self):
        paned = ttk.PanedWindow(self.tab_single, orient = 'horizontal')
        paned.pack(fill = 'both', expand = True, padx = 4, pady = 4)
        frame_left = tk.Frame(paned, bg = self.c_card, bd = 1, relief = 'solid', width = 360)
        paned.add(frame_left, weight = 1)
        tk.Label(frame_left, text = '⚙️ 货柜参数配置', font = ('微软雅黑', 9, 'bold'), bg = '#EAECEF', fg = self.c_primary, anchor = 'w', padx = 10, height = 2).pack(fill = 'x')
        p_form = tk.Frame(frame_left, bg = self.c_card)
        p_form.pack(fill = 'x', padx = 10, pady = 6)
        tk.Label(p_form, text = '外贸公司:', bg = self.c_card, font = ('微软雅黑', 9, 'bold')).grid(row = 0, column = 0, sticky = 'w', pady = 3)
        self.var_s_company = tk.StringVar(value = '颐顿')
        self.combo_s_company = ttk.Combobox(p_form, textvariable = self.var_s_company, values = list(self.counterparty_map.keys()), width = 18)
        self.combo_s_company.grid(row = 0, column = 1, sticky = 'ew', pady = 3, padx = 4)
        self.combo_s_company.bind('<<ComboboxSelected>>', self.on_single_param_change)
        tk.Label(p_form, text = '合同号:', bg = self.c_card).grid(row = 1, column = 0, sticky = 'w', pady = 3)
        self.var_s_contract = tk.StringVar(value = 'TS2026002')
        ttk.Entry(p_form, textvariable = self.var_s_contract, width = 18).grid(row = 1, column = 1, sticky = 'ew', pady = 3, padx = 4)
        tk.Label(p_form, text = '集装箱号码:', bg = self.c_card).grid(row = 2, column = 0, sticky = 'w', pady = 3)
        self.var_s_container = tk.StringVar(value = 'UETU7873174.FX48543892')
        ttk.Entry(p_form, textvariable = self.var_s_container, width = 18).grid(row = 2, column = 1, sticky = 'ew', pady = 3, padx = 4)
        tk.Label(p_form, text = '凭证日期:', bg = self.c_card).grid(row = 3, column = 0, sticky = 'w', pady = 3)
        self.var_s_date = tk.StringVar(value = datetime.datetime.now().strftime('%Y%m%d'))
        ttk.Entry(p_form, textvariable = self.var_s_date, width = 18).grid(row = 3, column = 1, sticky = 'ew', pady = 3, padx = 4)
        tk.Label(p_form, text = '凭证序号:', bg = self.c_card).grid(row = 4, column = 0, sticky = 'w', pady = 3)
        self.var_s_seq = tk.StringVar(value = '1')
        ttk.Entry(p_form, textvariable = self.var_s_seq, width = 18).grid(row = 4, column = 1, sticky = 'ew', pady = 3, padx = 4)
        tk.Label(p_form, text = '汇率 (USD/RMB):', bg = self.c_card, font = ('微软雅黑', 9, 'bold')).grid(row = 5, column = 0, sticky = 'w', pady = 3)
        self.var_s_rate = tk.StringVar(value = '6.7322')
        entry_rate = ttk.Entry(p_form, textvariable = self.var_s_rate, width = 18)
        entry_rate.grid(row = 5, column = 1, sticky = 'ew', pady = 3, padx = 4)
        entry_rate.bind('<KeyRelease>', self.on_single_param_change)
        tk.Label(p_form, text = '接收仓库编码:', bg = self.c_card, font = ('微软雅黑', 9, 'bold')).grid(row = 6, column = 0, sticky = 'w', pady = 3)
        self.var_s_warehouse = tk.StringVar(value = '100')
        ttk.Entry(p_form, textvariable = self.var_s_warehouse, width = 18).grid(row = 6, column = 1, sticky = 'ew', pady = 3, padx = 4)
        ttk.Separator(p_form, orient = 'horizontal').grid(row = 7, column = 0, columnspan = 2, sticky = 'ew', pady = 6)
        tk.Label(p_form, text = '海运费 (USD $):', bg = self.c_card, fg = '#1F4E78').grid(row = 8, column = 0, sticky = 'w', pady = 3)
        self.var_s_ocean = tk.StringVar(value = '4600.00')
        entry_ocean = ttk.Entry(p_form, textvariable = self.var_s_ocean, width = 18)
        entry_ocean.grid(row = 8, column = 1, sticky = 'ew', pady = 3, padx = 4)
        entry_ocean.bind('<KeyRelease>', self.on_single_param_change)
        tk.Label(p_form, text = '内陆托箱费 (¥):', bg = self.c_card, fg = '#1F4E78').grid(row = 9, column = 0, sticky = 'w', pady = 3)
        self.var_s_truck = tk.StringVar(value = '4800.00')
        entry_truck = ttk.Entry(p_form, textvariable = self.var_s_truck, width = 18)
        entry_truck.grid(row = 9, column = 1, sticky = 'ew', pady = 3, padx = 4)
        entry_truck.bind('<KeyRelease>', self.on_single_param_change)
        tk.Label(p_form, text = '保费 (USD $):', bg = self.c_card, fg = '#1F4E78').grid(row = 10, column = 0, sticky = 'w', pady = 3)
        self.var_s_insurance = tk.StringVar(value = '71.90')
        entry_insurance = ttk.Entry(p_form, textvariable = self.var_s_insurance, width = 18)
        entry_insurance.grid(row = 10, column = 1, sticky = 'ew', pady = 3, padx = 4)
        entry_insurance.bind('<KeyRelease>', self.on_single_param_change)
        tk.Label(p_form, text = '代理/点费 (¥):', bg = self.c_card).grid(row = 11, column = 0, sticky = 'w', pady = 3)
        self.var_s_agency = tk.StringVar(value = '0.00')
        entry_agency = ttk.Entry(p_form, textvariable = self.var_s_agency, width = 18)
        entry_agency.grid(row = 11, column = 1, sticky = 'ew', pady = 3, padx = 4)
        entry_agency.bind('<KeyRelease>', self.on_single_param_change)
        p_form.columnconfigure(1, weight = 1)
        kpi_frame = tk.LabelFrame(frame_left, text = '📊 实时核算汇总', bg = '#F8FAFC', font = ('微软雅黑', 9, 'bold'), fg = self.c_primary)
        kpi_frame.pack(fill = 'x', padx = 10, pady = 4)
        self.lbl_kpi_cbm = tk.Label(kpi_frame, text = '总体积: 0.00 CBM', bg = '#F8FAFC', anchor = 'w')
        self.lbl_kpi_cbm.pack(fill = 'x', padx = 6, pady = 1)
        self.lbl_kpi_3fees = tk.Label(kpi_frame, text = '三费总计: ¥0.00', bg = '#F8FAFC', font = ('微软雅黑', 9, 'bold'), fg = '#D9534F', anchor = 'w')
        self.lbl_kpi_3fees.pack(fill = 'x', padx = 6, pady = 1)
        self.lbl_kpi_goods = tk.Label(kpi_frame, text = '货款未税总额: ¥0.00', bg = '#F8FAFC', anchor = 'w')
        self.lbl_kpi_goods.pack(fill = 'x', padx = 6, pady = 1)
        self.lbl_kpi_cost = tk.Label(kpi_frame, text = '综合到岸总成本: ¥0.00', bg = '#F8FAFC', font = ('微软雅黑', 9, 'bold'), fg = self.c_primary, anchor = 'w')
        self.lbl_kpi_cost.pack(fill = 'x', padx = 6, pady = 1)
        ttk.Checkbutton(frame_left, text = '同时导出独立上传单表 (采购/新品/凭证)', variable = self.var_export_individual).pack(fill = 'x', padx = 10, pady = (2, 2))
        btn_export_single = ttk.Button(frame_left, text = '💾 导出当前柜全套上传表', style = 'Success.TButton', command = self.on_export_single_container)
        btn_export_single.pack(fill = 'x', padx = 10, pady = (4, 8))
        frame_right = tk.Frame(paned, bg = self.c_card, bd = 1, relief = 'solid')
        paned.add(frame_right, weight = 3)
        r_bar = tk.Frame(frame_right, bg = '#F0F4F8', height = 32)
        r_bar.pack(fill = 'x')
        r_bar.pack_propagate(False)
        tk.Label(r_bar, text = '📦 货物明细列表 (带空格品名自动转为-编码)', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        btn_add_item = ttk.Button(r_bar, text = '➕ 添加型号', width = 8, command = self.on_add_item_row)
        btn_add_item.pack(side = 'left', padx = 3)
        btn_edit_item = ttk.Button(r_bar, text = '✏️ 编辑选中', width = 8, command = self.on_edit_item_row)
        btn_edit_item.pack(side = 'left', padx = 3)
        btn_del_item = ttk.Button(r_bar, text = '🗑️ 删除选中', width = 8, command = self.on_delete_item_row)
        btn_del_item.pack(side = 'left', padx = 3)
        btn_clear_items = ttk.Button(r_bar, text = '🧹 清空明细', width = 8, command = self.on_clear_items)
        btn_clear_items.pack(side = 'left', padx = 3)
        item_cols = ('cat', 'model', 'raw_model', 'boxes', 'pack_qty', 'qty', 'unit_price', 'total_price', 'untaxed_total', 'cbm', 'fee_alloc', 'landing_unit')
        self.tree_items = ttk.Treeview(frame_right, columns = item_cols, show = 'headings', selectmode = 'extended')
        self.tree_items.heading('cat', text = '品名')
        self.tree_items.heading('model', text = '规范品目编码(无空格)')
        self.tree_items.heading('raw_model', text = '原始型号规格')
        self.tree_items.heading('boxes', text = '箱数')
        self.tree_items.heading('pack_qty', text = '装数')
        self.tree_items.heading('qty', text = '数量(台)')
        self.tree_items.heading('unit_price', text = '含税单价(¥)')
        self.tree_items.heading('total_price', text = '含税总价(¥)')
        self.tree_items.heading('untaxed_total', text = '未税总额(¥)')
        self.tree_items.heading('cbm', text = '体积(CBM)')
        self.tree_items.heading('fee_alloc', text = '分摊附带费(¥)')
        self.tree_items.heading('landing_unit', text = '单台到岸成本(¥)')
        self.tree_items.column('cat', width = 65, anchor = 'center')
        self.tree_items.column('model', width = 125, anchor = 'w')
        self.tree_items.column('raw_model', width = 110, anchor = 'w')
        self.tree_items.column('boxes', width = 45, anchor = 'e')
        self.tree_items.column('pack_qty', width = 40, anchor = 'e')
        self.tree_items.column('qty', width = 50, anchor = 'e')
        self.tree_items.column('unit_price', width = 80, anchor = 'e')
        self.tree_items.column('total_price', width = 85, anchor = 'e')
        self.tree_items.column('untaxed_total', width = 85, anchor = 'e')
        self.tree_items.column('cbm', width = 70, anchor = 'e')
        self.tree_items.column('fee_alloc', width = 80, anchor = 'e')
        self.tree_items.column('landing_unit', width = 90, anchor = 'e')
        sb_iv = ttk.Scrollbar(frame_right, orient = 'vertical', command = self.tree_items.yview)
        sb_ih = ttk.Scrollbar(frame_right, orient = 'horizontal', command = self.tree_items.xview)
        self.tree_items.configure(yscrollcommand = sb_iv.set, xscrollcommand = sb_ih.set)
        self.tree_items.pack(fill = 'both', expand = True, side = 'top')
        sb_iv.pack(side = 'right', fill = 'y')
        sb_ih.pack(side = 'bottom', fill = 'x')
        self.tree_items.bind('<Double-1>', (lambda e: self.on_edit_item_row()))

    
    def build_tab_mapping(self):
        card_m = tk.Frame(self.tab_mapping, bg = self.c_card, bd = 1, relief = 'solid')
        card_m.pack(fill = 'both', expand = True, padx = 6, pady = 4)
        m_top = tk.Frame(card_m, bg = '#F0F4F8', height = 40)
        m_top.pack(fill = 'x')
        m_top.pack_propagate(False)
        tk.Label(m_top, text = '💡 提示：当上传表格的列顺序发生变动时，可在下方灵活调整各字段对应的 Excel 列。', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        btn_auto_detect = ttk.Button(m_top, text = '🔍 重新智能自动匹配', command = self.on_auto_rematch_headers)
        btn_auto_detect.pack(side = 'right', padx = 6)
        btn_save_mapping = ttk.Button(m_top, text = '💾 保存当前映射配置', style = 'Primary.TButton', command = self.on_save_header_mapping_config)
        btn_save_mapping.pack(side = 'right', padx = 6)
        btn_reset_mapping = ttk.Button(m_top, text = '🔄 恢复默认标准映射', command = self.on_reset_header_mapping)
        btn_reset_mapping.pack(side = 'right', padx = 6)
        paned_map = ttk.PanedWindow(card_m, orient = 'horizontal')
        paned_map.pack(fill = 'both', expand = True, padx = 6, pady = 6)
        frame_sum_map = tk.LabelFrame(paned_map, text = '1. 货柜汇总表 (Sheet: 汇总) 字段列映射', bg = self.c_card, font = ('微软雅黑', 9, 'bold'), fg = self.c_primary)
        paned_map.add(frame_sum_map, weight = 1)
        canvas_sum = tk.Canvas(frame_sum_map, bg = self.c_card, bd = 0, highlightthickness = 0)
        sb_sum_map = ttk.Scrollbar(frame_sum_map, orient = 'vertical', command = canvas_sum.yview)
        self.inner_sum_map = tk.Frame(canvas_sum, bg = self.c_card)
        self.inner_sum_map.bind('<Configure>', (lambda e: canvas_sum.configure(scrollregion = canvas_sum.bbox('all'))))
        canvas_sum.create_window((0, 0), window = self.inner_sum_map, anchor = 'nw')
        canvas_sum.configure(yscrollcommand = sb_sum_map.set)
        canvas_sum.pack(side = 'left', fill = 'both', expand = True, padx = 6, pady = 6)
        sb_sum_map.pack(side = 'right', fill = 'y')
        frame_det_map = tk.LabelFrame(paned_map, text = '2. 发货明细表 (Sheet: 发货明细) 字段列映射', bg = self.c_card, font = ('微软雅黑', 9, 'bold'), fg = self.c_primary)
        paned_map.add(frame_det_map, weight = 1)
        canvas_det = tk.Canvas(frame_det_map, bg = self.c_card, bd = 0, highlightthickness = 0)
        sb_det_map = ttk.Scrollbar(frame_det_map, orient = 'vertical', command = canvas_det.yview)
        self.inner_det_map = tk.Frame(canvas_det, bg = self.c_card)
        self.inner_det_map.bind('<Configure>', (lambda e: canvas_det.configure(scrollregion = canvas_det.bbox('all'))))
        canvas_det.create_window((0, 0), window = self.inner_det_map, anchor = 'nw')
        canvas_det.configure(yscrollcommand = sb_det_map.set)
        canvas_det.pack(side = 'left', fill = 'both', expand = True, padx = 6, pady = 6)
        sb_det_map.pack(side = 'right', fill = 'y')
        self.sum_map_vars = { }
        self.sum_map_combos = { }
        self.det_map_vars = { }
        self.det_map_combos = { }
        self.render_header_mapping_widgets()

    
    def render_header_mapping_widgets(self):
        for child in self.inner_sum_map.winfo_children():
            child.destroy()
        self.sum_map_vars.clear()
        self.sum_map_combos.clear()
        sum_options = self.get_column_options(self.raw_sum_headers)
        active_sum_map = self.header_mapping.get('summary', { })
        for row_idx, (f_key, f_def) in enumerate(core.SUMMARY_FIELD_DEFS.items()):
            lbl_text = f'''{f_def['label']}:'''
            tk.Label(self.inner_sum_map, text = lbl_text, font = ('微软雅黑', 9), bg = self.c_card, anchor = 'w').grid(row = row_idx, column = 0, sticky = 'w', padx = 4, pady = 3)
            var = tk.StringVar()
            cur_idx = active_sum_map.get(f_key)
            if cur_idx is not None and 0 <= cur_idx < len(sum_options):
                var.set(sum_options[cur_idx])
            elif sum_options:
                var.set(sum_options[0])
            cb = ttk.Combobox(self.inner_sum_map, textvariable = var, values = sum_options, state = 'readonly', width = 28)
            cb.grid(row = row_idx, column = 1, sticky = 'ew', padx = 6, pady = 3)
            self.sum_map_vars[f_key] = var
            self.sum_map_combos[f_key] = cb
        self.inner_sum_map.columnconfigure(1, weight = 1)
        for child in self.inner_det_map.winfo_children():
            child.destroy()
        self.det_map_vars.clear()
        self.det_map_combos.clear()
        det_options = self.get_column_options(self.raw_det_headers)
        active_det_map = self.header_mapping.get('detail', { })
        for row_idx, (f_key, f_def) in enumerate(core.DETAIL_FIELD_DEFS.items()):
            lbl_text = f'''{f_def['label']}:'''
            tk.Label(self.inner_det_map, text = lbl_text, font = ('微软雅黑', 9), bg = self.c_card, anchor = 'w').grid(row = row_idx, column = 0, sticky = 'w', padx = 4, pady = 3)
            var = tk.StringVar()
            cur_idx = active_det_map.get(f_key)
            if cur_idx is not None and 0 <= cur_idx < len(det_options):
                var.set(det_options[cur_idx])
            elif det_options:
                var.set(det_options[0])
            cb = ttk.Combobox(self.inner_det_map, textvariable = var, values = det_options, state = 'readonly', width = 28)
            cb.grid(row = row_idx, column = 1, sticky = 'ew', padx = 6, pady = 3)
            self.det_map_vars[f_key] = var
            self.det_map_combos[f_key] = cb
        self.inner_det_map.columnconfigure(1, weight = 1)

    
    def get_column_options(self, headers_list):
        if not headers_list:
            return [ f'''[第 {i + 1} 列] (未读取到表头)''' for i in range(30) ]
        options = []
        for i, h in enumerate(headers_list):
            h_str = str(h).strip() if h is not None else '(空)'
            options.append(f'''[第 {i + 1} 列] {h_str}''')
        return options

    
    def on_auto_rematch_headers(self):
        if not (self.raw_sum_headers) and not (self.raw_det_headers):
            file_path = self.var_source_file.get().strip()
            if file_path and os.path.exists(file_path):
                self._read_raw_headers_from_file(file_path)
        sum_m = core.auto_match_columns(self.raw_sum_headers, core.SUMMARY_FIELD_DEFS)
        det_m = core.auto_match_columns(self.raw_det_headers, core.DETAIL_FIELD_DEFS)
        self.header_mapping = {
            'summary': sum_m,
            'detail': det_m }
        self.render_header_mapping_widgets()
        self.log('[*] 已根据当前 Excel 文件表头文字完成智能自动匹配！', 'ok')
        messagebox.showinfo('匹配完成', '已根据当前 Excel 文件表头文字完成智能自动匹配！\n如有需要微调的字段，可手动选择对应列。')

    
    def on_save_header_mapping_config(self):
        sum_m = { }
        sum_opts = self.get_column_options(self.raw_sum_headers)
        for k, var in self.sum_map_vars.items():
            val = var.get()
            if val in sum_opts:
                sum_m[k] = sum_opts.index(val)
        det_m = { }
        det_opts = self.get_column_options(self.raw_det_headers)
        for k, var in self.det_map_vars.items():
            val = var.get()
            if val in det_opts:
                det_m[k] = det_opts.index(val)
        self.header_mapping = {
            'summary': sum_m,
            'detail': det_m }
        if core.save_header_mapping(self.header_mapping):
            self.log('[OK] 表头映射配置已成功保存到 header_mapping_config.json', 'ok')
            messagebox.showinfo('成功', '表头映射配置已成功保存！\n后续解析与生成将采用此配置。')
            return None

    
    def on_reset_header_mapping(self):
        if messagebox.askyesno('确认', '确定要重置为默认标准表头映射吗？'):
            sum_m = core.auto_match_columns(self.raw_sum_headers, core.SUMMARY_FIELD_DEFS)
            det_m = core.auto_match_columns(self.raw_det_headers, core.DETAIL_FIELD_DEFS)
            self.header_mapping = {
                'summary': sum_m,
                'detail': det_m }
            core.save_header_mapping(self.header_mapping)
            self.render_header_mapping_widgets()
            self.log('[*] 已恢复默认标准映射配置')
            return None

    
    def _read_raw_headers_from_file(self, file_path):
        '''读取源文件表头用于映射配置界面'''
        
        try:
            wb = openpyxl.load_workbook(file_path, read_only = True, data_only = True)
            s_sum = None
            s_det = None
            for s in wb.sheetnames:
                if '汇总' in s:
                    s_sum = s
                    continue
                if '明细' in s or '发货' in s:
                    s_det = s
            if not s_sum:
                s_sum = wb.sheetnames[0]
            if not s_det:
                s_det = wb.sheetnames[1] if len(wb.sheetnames) > 1 else wb.sheetnames[0]
            ws_sum = wb[s_sum]
            ws_det = wb[s_det]
            sum_rows = list(ws_sum.iter_rows(values_only = True))[:10]
            det_rows = list(ws_det.iter_rows(values_only = True))[:10]
            s_idx = 2
            for i, r in enumerate(sum_rows):
                r_str = ' '.join(( str(c or '') for c in r ))
                if '合同号' in r_str:
                    s_idx = i
                    break
            self.raw_sum_headers = [ c for c in sum_rows[s_idx] ] if sum_rows else []
            d_idx = 2
            for i, r in enumerate(det_rows):
                r_str = ' '.join(( str(c or '') for c in r ))
                if '型号' in r_str:
                    d_idx = i
                    break
            self.raw_det_headers = [ c for c in det_rows[d_idx] ] if det_rows else []
            wb.close()
            return None
        except Exception as e:
            self.log(f'''[Warn] 读取表头信息失败: {e}''', 'warn')
            return None


    
    def build_tab_template(self):
        card_t_top = tk.Frame(self.tab_template, bg = self.c_card, bd = 1, relief = 'solid')
        card_t_top.pack(fill = 'x', padx = 6, pady = 4)
        t_top_bar = tk.Frame(card_t_top, bg = '#F0F4F8', height = 36)
        t_top_bar.pack(fill = 'x')
        t_top_bar.pack_propagate(False)
        tk.Label(t_top_bar, text = '📁 ECOUNT ERP 官方上传格式模板 (自动识别表头、列序与批注规范)', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        btn_auto_scan = ttk.Button(t_top_bar, text = '🔍 智能自动检测模板', command = self.on_auto_detect_templates)
        btn_auto_scan.pack(side = 'right', padx = 6)
        btn_save_tpl = ttk.Button(t_top_bar, text = '💾 保存模板路径配置', style = 'Primary.TButton', command = self.on_save_template_config)
        btn_save_tpl.pack(side = 'right', padx = 6)
        btn_reset_tpl = ttk.Button(t_top_bar, text = '🔄 恢复默认设置', command = self.on_reset_template_config)
        btn_reset_tpl.pack(side = 'right', padx = 6)
        f_tpl_grid = tk.Frame(card_t_top, bg = self.c_card)
        f_tpl_grid.pack(fill = 'x', padx = 12, pady = 8)
        tk.Label(f_tpl_grid, text = '1. 采购入库模板 (25列):', font = ('微软雅黑', 9, 'bold'), bg = self.c_card).grid(row = 0, column = 0, sticky = 'w', pady = 4)
        self.var_tpl_purchase = tk.StringVar(value = self.template_config.get('purchase', ''))
        self.entry_tpl_purchase = ttk.Entry(f_tpl_grid, textvariable = self.var_tpl_purchase, width = 50)
        self.entry_tpl_purchase.grid(row = 0, column = 1, sticky = 'ew', padx = 8, pady = 4)
        ttk.Button(f_tpl_grid, text = '浏览...', width = 8, command = (lambda : self.on_browse_template_file('purchase'))).grid(row = 0, column = 2, padx = 4, pady = 4)
        self.lbl_status_purchase = tk.Label(f_tpl_grid, text = '检测中...', font = ('微软雅黑', 9), bg = self.c_card, fg = '#666666')
        self.lbl_status_purchase.grid(row = 0, column = 3, sticky = 'w', padx = 6, pady = 4)
        tk.Label(f_tpl_grid, text = '2. 新增品目模板 (17列):', font = ('微软雅黑', 9, 'bold'), bg = self.c_card).grid(row = 1, column = 0, sticky = 'w', pady = 4)
        self.var_tpl_new_item = tk.StringVar(value = self.template_config.get('new_item', ''))
        self.entry_tpl_new_item = ttk.Entry(f_tpl_grid, textvariable = self.var_tpl_new_item, width = 50)
        self.entry_tpl_new_item.grid(row = 1, column = 1, sticky = 'ew', padx = 8, pady = 4)
        ttk.Button(f_tpl_grid, text = '浏览...', width = 8, command = (lambda : self.on_browse_template_file('new_item'))).grid(row = 1, column = 2, padx = 4, pady = 4)
        self.lbl_status_new_item = tk.Label(f_tpl_grid, text = '检测中...', font = ('微软雅黑', 9), bg = self.c_card, fg = '#666666')
        self.lbl_status_new_item.grid(row = 1, column = 3, sticky = 'w', padx = 6, pady = 4)
        tk.Label(f_tpl_grid, text = '3. 一般凭证模板 (14列):', font = ('微软雅黑', 9, 'bold'), bg = self.c_card).grid(row = 2, column = 0, sticky = 'w', pady = 4)
        self.var_tpl_voucher = tk.StringVar(value = self.template_config.get('voucher', ''))
        self.entry_tpl_voucher = ttk.Entry(f_tpl_grid, textvariable = self.var_tpl_voucher, width = 50)
        self.entry_tpl_voucher.grid(row = 2, column = 1, sticky = 'ew', padx = 8, pady = 4)
        ttk.Button(f_tpl_grid, text = '浏览...', width = 8, command = (lambda : self.on_browse_template_file('voucher'))).grid(row = 2, column = 2, padx = 4, pady = 4)
        self.lbl_status_voucher = tk.Label(f_tpl_grid, text = '检测中...', font = ('微软雅黑', 9), bg = self.c_card, fg = '#666666')
        self.lbl_status_voucher.grid(row = 2, column = 3, sticky = 'w', padx = 6, pady = 4)
        f_tpl_grid.columnconfigure(1, weight = 1)
        f_tpl_opt = tk.Frame(card_t_top, bg = '#FAFAFA', bd = 1, relief = 'groove')
        f_tpl_opt.pack(fill = 'x', padx = 12, pady = (0, 8))
        tk.Label(f_tpl_opt, text = '⚙️ 导出模式选项:', font = ('微软雅黑', 9, 'bold'), bg = '#FAFAFA', fg = self.c_primary).pack(side = 'left', padx = 8, pady = 6)
        ttk.Checkbutton(f_tpl_opt, text = '导出 5-in-1 全套综合上传工作簿 (含批注与格式)', variable = self.var_export_bundle).pack(side = 'left', padx = 8, pady = 6)
        ttk.Checkbutton(f_tpl_opt, text = '同时同步导出 ECOUNT 独立上传单表 (采购入库/新增品目/一般凭证单表)', variable = self.var_export_individual).pack(side = 'left', padx = 8, pady = 6)
        card_t_bot = tk.Frame(self.tab_template, bg = self.c_card, bd = 1, relief = 'solid')
        card_t_bot.pack(fill = 'both', expand = True, padx = 6, pady = 4)
        inspector_bar = tk.Frame(card_t_bot, bg = '#F0F4F8', height = 36)
        inspector_bar.pack(fill = 'x')
        inspector_bar.pack_propagate(False)
        tk.Label(inspector_bar, text = '🔍 模板字段与批注规范深度查看器', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        tk.Label(inspector_bar, text = '选择模板:', bg = '#F0F4F8', font = ('微软雅黑', 9)).pack(side = 'left', padx = (15, 4))
        self.var_tpl_view_type = tk.StringVar(value = 'purchase')
        self.combo_tpl_view = ttk.Combobox(inspector_bar, textvariable = self.var_tpl_view_type, values = [
            'purchase: 采购入库格式模板',
            'new_item: 新增品目格式模板',
            'voucher: 一般凭证格式模板'], state = 'readonly', width = 28)
        self.combo_tpl_view.pack(side = 'left', padx = 4)
        self.combo_tpl_view.bind('<<ComboboxSelected>>', self.on_tpl_view_changed)
        btn_refresh_tpl_view = ttk.Button(inspector_bar, text = '🔄 刷新解析', command = self.refresh_template_inspector)
        btn_refresh_tpl_view.pack(side = 'left', padx = 8)
        cols = ('col_idx', 'coord', 'header', 'bound_key', 'limits', 'enums', 'comment_snippet')
        self.tree_template_rules = ttk.Treeview(card_t_bot, columns = cols, show = 'headings', selectmode = 'browse')
        self.tree_template_rules.heading('col_idx', text = '列号')
        self.tree_template_rules.heading('coord', text = '单元格')
        self.tree_template_rules.heading('header', text = '模板表头字段名')
        self.tree_template_rules.heading('bound_key', text = '系统绑定数据字段')
        self.tree_template_rules.heading('limits', text = '字数/位数限制规范')
        self.tree_template_rules.heading('enums', text = '枚举选项/代码定义')
        self.tree_template_rules.heading('comment_snippet', text = '原始批注简要')
        self.tree_template_rules.column('col_idx', width = 45, anchor = 'center')
        self.tree_template_rules.column('coord', width = 55, anchor = 'center')
        self.tree_template_rules.column('header', width = 140, anchor = 'w')
        self.tree_template_rules.column('bound_key', width = 140, anchor = 'w')
        self.tree_template_rules.column('limits', width = 160, anchor = 'w')
        self.tree_template_rules.column('enums', width = 220, anchor = 'w')
        self.tree_template_rules.column('comment_snippet', width = 260, anchor = 'w')
        sb_v_t = ttk.Scrollbar(card_t_bot, orient = 'vertical', command = self.tree_template_rules.yview)
        sb_h_t = ttk.Scrollbar(card_t_bot, orient = 'horizontal', command = self.tree_template_rules.xview)
        self.tree_template_rules.configure(yscrollcommand = sb_v_t.set, xscrollcommand = sb_h_t.set)
        self.tree_template_rules.pack(fill = 'both', expand = True, side = 'top')
        sb_v_t.pack(side = 'right', fill = 'y')
        sb_h_t.pack(side = 'bottom', fill = 'x')
        self.tree_template_rules.bind('<<TreeviewSelect>>', self.on_template_rule_row_select)
        detail_frame = tk.Frame(card_t_bot, bg = '#F7F9FC', height = 75, bd = 1, relief = 'groove')
        detail_frame.pack(fill = 'x', padx = 4, pady = 4)
        detail_frame.pack_propagate(False)
        self.lbl_selected_comment_title = tk.Label(detail_frame, text = '【字段批注完整详情】: 请点击上方表格行查看具体字段的官方批注说明', font = ('微软雅黑', 9, 'bold'), bg = '#F7F9FC', fg = self.c_primary, anchor = 'w')
        self.lbl_selected_comment_title.pack(fill = 'x', padx = 8, pady = (3, 0))
        self.txt_selected_comment = tk.Text(detail_frame, bg = '#FFFFFF', fg = '#24292F', font = ('微软雅黑', 9), height = 2, bd = 1, relief = 'solid', wrap = 'word')
        self.txt_selected_comment.pack(fill = 'both', expand = True, padx = 8, pady = (2, 4))
        self.txt_selected_comment.insert('1.0', '暂未选中任何列。点击上方表格任意列即可查看其在 ECOUNT 模板中的完整批注限制、字数上限及枚举值说明。')
        self.txt_selected_comment.config(state = 'disabled')
        self.refresh_all_template_statuses()
        self.refresh_template_inspector()

    
    def refresh_all_template_statuses(self):
        '''刷新三个模板文件的检测状态'''
        p_path = self.var_tpl_purchase.get().strip()
        info_p = core.inspect_template_file(p_path, template_type = 'purchase')
        if info_p['exists']:
            cmt_cnt = sum(( 1 for c in info_p['columns'] if c['has_comment'] ))
            self.lbl_status_purchase.config(text = f'''✅ 正常 ({info_p['max_columns']}列 / {cmt_cnt}条批注)''', fg = self.c_success)
        else:
            self.lbl_status_purchase.config(text = '❌ 未找到文件 (将使用系统内置格式)', fg = '#D9534F')
        n_path = self.var_tpl_new_item.get().strip()
        info_n = core.inspect_template_file(n_path, template_type = 'new_item')
        if info_n['exists']:
            cmt_cnt = sum(( 1 for c in info_n['columns'] if c['has_comment'] ))
            self.lbl_status_new_item.config(text = f'''✅ 正常 ({info_n['max_columns']}列 / {cmt_cnt}条批注)''', fg = self.c_success)
        else:
            self.lbl_status_new_item.config(text = '❌ 未找到文件 (将使用系统内置格式)', fg = '#D9534F')
        v_path = self.var_tpl_voucher.get().strip()
        info_v = core.inspect_template_file(v_path, template_type = 'voucher')
        if info_v['exists']:
            cmt_cnt = sum(( 1 for c in info_v['columns'] if c['has_comment'] ))
            self.lbl_status_voucher.config(text = f'''✅ 正常 ({info_v['max_columns']}列 / {cmt_cnt}条批注)''', fg = self.c_success)
            return None
        self.lbl_status_voucher.config(text = '❌ 未找到文件 (将使用系统内置格式)', fg = '#D9534F')

    
    def refresh_template_inspector(self):
        '''刷新批注查看器表格数据'''
        for item in self.tree_template_rules.get_children():
            self.tree_template_rules.delete(item)
        raw_type = self.var_tpl_view_type.get().split(':')[0].strip()
        if raw_type == 'purchase':
            f_path = self.var_tpl_purchase.get().strip()
            field_defs = core.PURCHASE_TEMPLATE_FIELD_DEFS
        elif raw_type == 'new_item':
            f_path = self.var_tpl_new_item.get().strip()
            field_defs = core.NEW_ITEM_TEMPLATE_FIELD_DEFS
        else:
            f_path = self.var_tpl_voucher.get().strip()
            field_defs = core.VOUCHER_TEMPLATE_FIELD_DEFS
        info = core.inspect_template_file(f_path, template_type = raw_type)
        self.cached_inspector_columns = info.get('columns', [])
        if not info['exists']:
            self.txt_selected_comment.config(state = 'normal')
            self.txt_selected_comment.delete('1.0', 'end')
            self.txt_selected_comment.insert('1.0', f'''未找到模板文件: {f_path}\n请先在上方指定有效的 Excel 模板文件路径。''')
            self.txt_selected_comment.config(state = 'disabled')
            return None
        headers = [ c['header'] for c in info['columns'] ]
        matched_mapping = core.auto_match_columns(headers, field_defs)
        for col in info['columns']:
            c_idx = col['col_idx']
            h_text = col['header']
            bound_key_label = '未绑定'
            for f_k, f_idx in matched_mapping.items():
                if f_idx == c_idx - 1:
                    bound_key_label = f'''{f_k} ({field_defs[f_k]['label']})'''
                    break
            cmt_snippet = col['comment'].replace('\n', ' ').strip() if col['comment'] else '(无批注)'
            if len(cmt_snippet) > 40:
                cmt_snippet = cmt_snippet[:40] + '...'
            vals = (c_idx, col['coord'], h_text, bound_key_label, col['limit_desc'], col['enum_desc'], cmt_snippet)
            self.tree_template_rules.insert('', 'end', iid = str(c_idx - 1), values = vals)

    
    def on_tpl_view_changed(self, event = None):
        self.refresh_template_inspector()

    
    def on_template_rule_row_select(self, event = None):
        sel = self.tree_template_rules.selection()
        if not sel:
            return None
        idx = int(sel[0])
        if 0 <= idx < len(getattr(self, 'cached_inspector_columns', [])):
            col = self.cached_inspector_columns[idx]
            self.lbl_selected_comment_title.config(text = f'''【字段批注完整详情】: 第 {col['col_idx']} 列 [{col['coord']}] - {col['header']}  (限制: {col['limit_desc']})''')
            self.txt_selected_comment.config(state = 'normal')
            self.txt_selected_comment.delete('1.0', 'end')
            if col['comment']:
                self.txt_selected_comment.insert('1.0', col['comment'])
            else:
                self.txt_selected_comment.insert('1.0', '此列单元格无批注规范。')
            self.txt_selected_comment.config(state = 'disabled')
            return None

    
    def on_browse_template_file(self, template_type):
        initial_dir = os.path.dirname(self.template_config.get(template_type, '')) or os.path.expanduser('~\\Downloads')
        path = filedialog.askopenfilename(title = '选择 ECOUNT ERP 官方上传格式模板', initialdir = initial_dir, filetypes = [
            ('Excel 文件', '*.xlsx *.xls'),
            ('所有文件', '*.*')])
        if path:
            if template_type == 'purchase':
                self.var_tpl_purchase.set(path)
            elif template_type == 'new_item':
                self.var_tpl_new_item.set(path)
            elif template_type == 'voucher':
                self.var_tpl_voucher.set(path)
            self.refresh_all_template_statuses()
            self.refresh_template_inspector()
            return None

    
    def on_auto_detect_templates(self):
        detected_cfg = core.load_template_config()
        self.var_tpl_purchase.set(detected_cfg.get('purchase', ''))
        self.var_tpl_new_item.set(detected_cfg.get('new_item', ''))
        self.var_tpl_voucher.set(detected_cfg.get('voucher', ''))
        self.refresh_all_template_statuses()
        self.refresh_template_inspector()
        self.log('[*] 已重新自动扫描常用目录并检测模板状态！', 'ok')
        messagebox.showinfo('检测完成', '已自动扫描常用目录并更新模板配置！')

    
    def on_save_template_config(self):
        self.template_config['purchase'] = self.var_tpl_purchase.get().strip()
        self.template_config['new_item'] = self.var_tpl_new_item.get().strip()
        self.template_config['voucher'] = self.var_tpl_voucher.get().strip()
        self.template_config['export_bundle_workbook'] = self.var_export_bundle.get()
        self.template_config['export_individual_files'] = self.var_export_individual.get()
        if core.save_template_config(self.template_config):
            self.log('[OK] 模板路径与导出配置已成功保存到 template_config.json', 'ok')
            messagebox.showinfo('成功', '模板配置已成功保存！\n后续生成工作簿时将自动应用上述模板与批注。')
        self.refresh_all_template_statuses()
        self.refresh_template_inspector()

    
    def on_reset_template_config(self):
        if messagebox.askyesno('确认', '确定要恢复默认模板路径设置吗？'):
            self.template_config = dict(core.DEFAULT_TEMPLATE_CONFIG)
            core.save_template_config(self.template_config)
            self.var_tpl_purchase.set(self.template_config['purchase'])
            self.var_tpl_new_item.set(self.template_config['new_item'])
            self.var_tpl_voucher.set(self.template_config['voucher'])
            self.var_export_bundle.set(True)
            self.var_export_individual.set(False)
            self.refresh_all_template_statuses()
            self.refresh_template_inspector()
            self.log('[*] 已恢复默认模板配置')
            return None

    
    def build_tab_master(self):
        card_m = tk.Frame(self.tab_master, bg = self.c_card, bd = 1, relief = 'solid')
        card_m.pack(fill = 'both', expand = True, padx = 6, pady = 4)
        m_bar = tk.Frame(card_m, bg = '#F0F4F8', height = 34)
        m_bar.pack(fill = 'x')
        m_bar.pack_propagate(False)
        tk.Label(m_bar, text = '品目搜索:', bg = '#F0F4F8', font = ('微软雅黑', 9, 'bold')).pack(side = 'left', padx = (10, 4))
        self.var_master_search = tk.StringVar()
        entry_m_search = ttk.Entry(m_bar, textvariable = self.var_master_search, width = 20)
        entry_m_search.pack(side = 'left', padx = 4)
        entry_m_search.bind('<KeyRelease>', self.on_filter_master_items)
        btn_load_m = ttk.Button(m_bar, text = '📂 载入外部品目主数据 (ESA009M / 上传表)', command = self.on_load_external_master)
        btn_load_m.pack(side = 'left', padx = 6)
        btn_add_m = ttk.Button(m_bar, text = '➕ 新增品目', command = self.on_add_master_item)
        btn_add_m.pack(side = 'left', padx = 3)
        btn_del_m = ttk.Button(m_bar, text = '🗑️ 删除选中', command = self.on_delete_master_item)
        btn_del_m.pack(side = 'left', padx = 3)
        self.lbl_master_count = tk.Label(m_bar, text = f'''品目总数: {len(self.master_items)}''', bg = '#F0F4F8', fg = '#555555')
        self.lbl_master_count.pack(side = 'right', padx = 12)
        m_cols = ('code', 'name', 'type', 'spec', 'comb1', 'kw', 'use_yn', 'conv', 'pack_qty', 'cbm', 'weight')
        self.tree_master = ttk.Treeview(card_m, columns = m_cols, show = 'headings', selectmode = 'extended')
        self.tree_master.heading('code', text = '品目编码 (无空格)')
        self.tree_master.heading('name', text = '品目名')
        self.tree_master.heading('type', text = '品目类型')
        self.tree_master.heading('spec', text = '规格信息')
        self.tree_master.heading('comb1', text = '品目组合1')
        self.tree_master.heading('kw', text = '查询关键词')
        self.tree_master.heading('use_yn', text = '使用')
        self.tree_master.heading('conv', text = '转换比率')
        self.tree_master.heading('pack_qty', text = '装数')
        self.tree_master.heading('cbm', text = '体积(CBM)')
        self.tree_master.heading('weight', text = '重量(kg)')
        self.tree_master.column('code', width = 120, anchor = 'w')
        self.tree_master.column('name', width = 120, anchor = 'w')
        self.tree_master.column('type', width = 70, anchor = 'center')
        self.tree_master.column('spec', width = 110, anchor = 'w')
        self.tree_master.column('comb1', width = 70, anchor = 'center')
        self.tree_master.column('kw', width = 70, anchor = 'w')
        self.tree_master.column('use_yn', width = 50, anchor = 'center')
        self.tree_master.column('conv', width = 65, anchor = 'center')
        self.tree_master.column('pack_qty', width = 50, anchor = 'e')
        self.tree_master.column('cbm', width = 75, anchor = 'e')
        self.tree_master.column('weight', width = 65, anchor = 'e')
        sb_mv = ttk.Scrollbar(card_m, orient = 'vertical', command = self.tree_master.yview)
        sb_mh = ttk.Scrollbar(card_m, orient = 'horizontal', command = self.tree_master.xview)
        self.tree_master.configure(yscrollcommand = sb_mv.set, xscrollcommand = sb_mh.set)
        self.tree_master.pack(fill = 'both', expand = True, side = 'top')
        sb_mv.pack(side = 'right', fill = 'y')
        sb_mh.pack(side = 'bottom', fill = 'x')
        self.refresh_master_table()

    
    def build_tab_counterparty(self):
        card_c = tk.Frame(self.tab_counterparty, bg = self.c_card, bd = 1, relief = 'solid')
        card_c.pack(fill = 'both', expand = True, padx = 6, pady = 4)
        c_bar = tk.Frame(card_c, bg = '#F0F4F8', height = 34)
        c_bar.pack(fill = 'x')
        c_bar.pack_propagate(False)
        tk.Label(c_bar, text = '往来单位映射表 (ECOUNT ERP 会计凭证及采购入库对应单位):', font = ('微软雅黑', 9, 'bold'), bg = '#F0F4F8', fg = self.c_primary).pack(side = 'left', padx = 10)
        btn_add_cp = ttk.Button(c_bar, text = '➕ 新增往来单位', command = self.on_add_counterparty)
        btn_add_cp.pack(side = 'left', padx = 3)
        btn_edit_cp = ttk.Button(c_bar, text = '✏️ 编辑选中', command = self.on_edit_counterparty)
        btn_edit_cp.pack(side = 'left', padx = 3)
        btn_del_cp = ttk.Button(c_bar, text = '🗑️ 删除选中', command = self.on_delete_counterparty)
        btn_del_cp.pack(side = 'left', padx = 3)
        btn_save_cp = ttk.Button(c_bar, text = '💾 保存配置', style = 'Primary.TButton', command = self.on_save_counterparty_config)
        btn_save_cp.pack(side = 'right', padx = 10)
        cp_cols = ('key', 'rmb_code', 'rmb_name', 'usd_code', 'usd_name')
        self.tree_cp = ttk.Treeview(card_c, columns = cp_cols, show = 'headings', selectmode = 'browse')
        self.tree_cp.heading('key', text = '公司简称/关键词')
        self.tree_cp.heading('rmb_code', text = '人民币往来单位编码')
        self.tree_cp.heading('rmb_name', text = '人民币往来单位名称')
        self.tree_cp.heading('usd_code', text = '美金往来单位编码')
        self.tree_cp.heading('usd_name', text = '美金往来单位名称')
        self.tree_cp.column('key', width = 120, anchor = 'center')
        self.tree_cp.column('rmb_code', width = 140, anchor = 'center')
        self.tree_cp.column('rmb_name', width = 220, anchor = 'w')
        self.tree_cp.column('usd_code', width = 140, anchor = 'center')
        self.tree_cp.column('usd_name', width = 240, anchor = 'w')
        sb_cp = ttk.Scrollbar(card_c, orient = 'vertical', command = self.tree_cp.yview)
        self.tree_cp.configure(yscrollcommand = sb_cp.set)
        self.tree_cp.pack(fill = 'both', expand = True, side = 'top')
        sb_cp.pack(side = 'right', fill = 'y')
        self.refresh_counterparty_table()

    
    def build_log_panel(self, parent):
        card_log = tk.Frame(parent, bg = self.c_card, bd = 1, relief = 'solid', height = 150)
        card_log.pack(fill = 'x', side = 'bottom', pady = (4, 0))
        card_log.pack_propagate(False)
        l_bar = tk.Frame(card_log, bg = '#EAECEF', height = 24)
        l_bar.pack(fill = 'x')
        l_bar.pack_propagate(False)
        tk.Label(l_bar, text = '📜 操作日志与状态反馈', font = ('微软雅黑', 8, 'bold'), bg = '#EAECEF', fg = '#333333').pack(side = 'left', padx = 8)
        btn_clear_log = ttk.Button(l_bar, text = '清空日志', width = 8, command = self.clear_log)
        btn_clear_log.pack(side = 'right', padx = 3)
        btn_copy_log = ttk.Button(l_bar, text = '复制日志', width = 8, command = self.copy_log)
        btn_copy_log.pack(side = 'right', padx = 3)
        self.txt_log = tk.Text(card_log, font = ('Consolas', 9), bg = '#1E1E1E', fg = '#D4D4D4', insertbackground = 'white', bd = 0)
        sb_log = ttk.Scrollbar(card_log, orient = 'vertical', command = self.txt_log.yview)
        self.txt_log.configure(yscrollcommand = sb_log.set)
        self.txt_log.pack(side = 'left', fill = 'both', expand = True)
        sb_log.pack(side = 'right', fill = 'y')
        self.txt_log.tag_config('info', foreground = '#61AFEF')
        self.txt_log.tag_config('ok', foreground = '#98C379')
        self.txt_log.tag_config('warn', foreground = '#E5C07B')
        self.txt_log.tag_config('err', foreground = '#E06C75')
        self.log('[System] 系统启动完成，欢迎使用 ECOUNT ERP 货柜费用分摊与导入表生成系统 v2.1！', 'ok')

    
    def log(self, message, level = 'info'):
        now_str = datetime.datetime.now().strftime('%H:%M:%S')
        formatted = f'''[{now_str}] {message}\n'''
        self.txt_log.insert('end', formatted, level)
        self.txt_log.see('end')

    
    def clear_log(self):
        self.txt_log.delete('1.0', 'end')

    
    def copy_log(self):
        content = self.txt_log.get('1.0', 'end')
        self.root.clipboard_clear()
        self.root.clipboard_append(content)
        messagebox.showinfo('提示', '日志已复制到剪贴板！')

    
    def auto_detect_default_file(self):
        base_dir = core.get_base_dir()
        for f in os.listdir(base_dir):
            if f.endswith('.xlsx'):
                if ('发货记录' in f or '汇总' in f) and not f.startswith('~$'):
                    full_path = os.path.join(base_dir, f)
                    self.var_source_file.set(full_path)
                    self.log(f'''[*] 自动检测到发货记录文件: {f}''')
                    self._read_raw_headers_from_file(full_path)
                    self.render_header_mapping_widgets()
                    threading.Thread(target = self._async_parse_source, args = (full_path,), daemon = True).start()
                    return None

    
    def on_browse_source_file(self):
        path = filedialog.askopenfilename(title = '选择发货记录汇总表 Excel', filetypes = [
            ('Excel 文件', '*.xlsx *.xls'),
            ('所有文件', '*.*')])
        if path:
            self.var_source_file.set(path)
            self._read_raw_headers_from_file(path)
            self.render_header_mapping_widgets()
            self.on_parse_source_file()
            return None

    
    def on_browse_output_dir(self):
        path = filedialog.askdirectory(title = '选择报表输出保存目录')
        if path:
            self.var_output_dir.set(path)
            return None

    
    def on_open_output_dir(self):
        out_dir = self.var_output_dir.get().strip()
        if not os.path.exists(out_dir):
            os.makedirs(out_dir, exist_ok = True)
        if sys.platform == 'win32':
            os.startfile(out_dir)
            return None
        subprocess.run([
            'xdg-open',
            out_dir])

    
    def on_parse_source_file(self):
        file_path = self.var_source_file.get().strip()
        if not file_path or not os.path.exists(file_path):
            messagebox.showerror('错误', '请先选择有效的发货记录汇总 Excel 文件！')
            return None
        self._read_raw_headers_from_file(file_path)
        self.render_header_mapping_widgets()
        threading.Thread(target = self._async_parse_source, args = (file_path,), daemon = True).start()

    
    def _async_parse_source(self, file_path):
        self.root.after(0, (lambda : self.set_busy_state(True, '正在解析发货汇总表...')))
        
        try:
            
            try:
                containers = core.parse_source_shipment_excel(file_path, custom_mapping = self.header_mapping, log_func = (lambda m: self.root.after(0, self.log, m)))
                self.containers = containers
                self.root.after(0, self._render_containers_table)
            except Exception as e:
                err_msg = str(e)
                self.root.after(0, (lambda msg=err_msg: self.log(f'''[Error] 解析失败: {msg}''', 'err')))
                self.root.after(0, (lambda msg=err_msg: messagebox.showerror('解析失败', msg)))

            return None
            return None
        finally:
            self.root.after(0, (lambda : self.set_busy_state(False, f'''已解析 {len(self.containers)} 个货柜''')))


    
    def _render_containers_table(self):
        for item in self.tree_containers.get_children():
            self.tree_containers.delete(item)
        for i, c in enumerate(self.containers):
            items_cnt = len(c.get('items', []))
            status_text = '正常' if items_cnt > 0 else '无明细'
            vals = ('☑ 是', c.get('seq_no', i + 1), c.get('trade_co', ''), c.get('contract_no', ''), c.get('container_no', ''), c.get('date', ''), f'''{c.get('rate', 6.7322):.4f}''', f'''{c.get('ocean_usd', 0):.2f}''', f'''{c.get('truck_rmb', 0):.2f}''', f'''{c.get('insurance_usd', 0):.2f}''', f'''{c.get('agency_rmb', 0):.2f}''', items_cnt, f'''{c.get('cbm', 0):.2f}''', status_text)
            self.tree_containers.insert('', 'end', iid = str(i), values = vals)
        self.lbl_container_count.config(text = f'''共 {len(self.containers)} 个货柜 (已全部默认选中)''')

    
    def on_select_all_containers(self):
        for item_id in self.tree_containers.get_children():
            cur_vals = list(self.tree_containers.item(item_id, 'values'))
            cur_vals[0] = '☑ 是'
            self.tree_containers.item(item_id, values = cur_vals)

    
    def on_invert_select_containers(self):
        for item_id in self.tree_containers.get_children():
            cur_vals = list(self.tree_containers.item(item_id, 'values'))
            cur_vals[0] = '☐ 否' if cur_vals[0] == '☑ 是' else '☑ 是'
            self.tree_containers.item(item_id, values = cur_vals)

    
    def on_double_click_container_row(self, event):
        item_id = self.tree_containers.identify_row(event.y)
        if item_id:
            cur_vals = list(self.tree_containers.item(item_id, 'values'))
            cur_vals[0] = '☐ 否' if cur_vals[0] == '☑ 是' else '☑ 是'
            self.tree_containers.item(item_id, values = cur_vals)
            return None

    
    def on_load_selected_to_single(self):
        selected = self.tree_containers.selection()
        if not selected:
            messagebox.showwarning('提示', '请在列表中选中一个货柜！')
            return None
        idx = int(selected[0])
        if 0 <= idx < len(self.containers):
            c = self.containers[idx]
            self.load_container_to_single_tab(c)
            self.notebook.select(self.tab_single)
            self.log(f'''[*] 已将货柜 [{c.get('contract_no')}] 载入单柜明细核算面板''')
            return None

    
    def load_container_to_single_tab(self, c):
        self.var_s_company.set(c.get('trade_co', '颐顿'))
        self.var_s_contract.set(c.get('contract_no', ''))
        self.var_s_container.set(c.get('container_no', ''))
        self.var_s_date.set(c.get('date', datetime.datetime.now().strftime('%Y%m%d')))
        self.var_s_seq.set(str(c.get('seq_no', 1)))
        self.var_s_rate.set(str(c.get('rate', 6.7322)))
        self.var_s_warehouse.set(str(c.get('warehouse', '100')))
        self.var_s_ocean.set(str(c.get('ocean_usd', 0.0)))
        self.var_s_truck.set(str(c.get('truck_rmb', 0.0)))
        self.var_s_insurance.set(str(c.get('insurance_usd', 0.0)))
        self.var_s_agency.set(str(c.get('agency_rmb', 0.0)))
        for row in self.tree_items.get_children():
            self.tree_items.delete(row)
        for item in c.get('items', []):
            raw_m = str(item.get('raw_model', item.get('model', ''))).strip()
            clean_c = core.sanitize_item_code(item.get('model', raw_m))
            vals = (item.get('cat', ''), clean_c, raw_m, item.get('boxes', 0), item.get('pack_qty', 1), item.get('qty', 0), f'''{item.get('unit_price', 0):.2f}''', f'''{item.get('total_price', 0):.2f}''', f'''{item.get('untaxed_total', 0):.2f}''', f'''{item.get('cbm', 0):.4f}''', '0.00', '0.00')
            self.tree_items.insert('', 'end', values = vals)
        self.recalculate_single_tab()

    
    def on_single_param_change(self, *args):
        self.recalculate_single_tab()

    
    def recalculate_single_tab(self):
        rate = core.safe_float(self.var_s_rate.get(), 6.7322)
        ocean_usd = core.safe_float(self.var_s_ocean.get(), 0.0)
        truck_rmb = core.safe_float(self.var_s_truck.get(), 0.0)
        insurance_usd = core.safe_float(self.var_s_insurance.get(), 0.0)
        agency_rmb = core.safe_float(self.var_s_agency.get(), 0.0)
        total_alloc_fees = round(ocean_usd * rate, 2) + truck_rmb + round(insurance_usd * rate, 2) + agency_rmb
        total_cbm = 0.0
        total_untaxed = 0.0
        items_rows = []
        for item_id in self.tree_items.get_children():
            vals = list(self.tree_items.item(item_id, 'values'))
            cbm = core.safe_float(vals[9])
            untaxed = core.safe_float(vals[8])
            qty = core.safe_int(vals[5], 1)
            total_cbm += cbm
            total_untaxed += untaxed
            items_rows.append((item_id, vals, cbm, untaxed, qty))
        for item_id, vals, cbm, untaxed, qty in items_rows:
            ratio = cbm / total_cbm if total_cbm > 0 else 0
            fee_alloc = round(total_alloc_fees * ratio, 2)
            landing_unit = round((untaxed + fee_alloc) / qty, 4) if qty > 0 else 0
            vals[10] = f'''{fee_alloc:.2f}'''
            vals[11] = f'''{landing_unit:.4f}'''
            self.tree_items.item(item_id, values = vals)
        total_cost = total_untaxed + total_alloc_fees
        self.lbl_kpi_cbm.config(text = f'''总体积: {total_cbm:.4f} CBM''')
        if agency_rmb > 0:
            self.lbl_kpi_3fees.config(text = f'''四费总计: ¥{total_alloc_fees:.2f} (海运:¥{round(ocean_usd * rate, 2):.2f}, 托箱:¥{truck_rmb:.2f}, 保费:¥{round(insurance_usd * rate, 2):.2f}, 代理:¥{agency_rmb:.2f})''')
        else:
            self.lbl_kpi_3fees.config(text = f'''三费总计: ¥{total_alloc_fees:.2f} (海运:¥{round(ocean_usd * rate, 2):.2f}, 托箱:¥{truck_rmb:.2f}, 保费:¥{round(insurance_usd * rate, 2):.2f})''')
        self.lbl_kpi_goods.config(text = f'''货款未税总额: ¥{total_untaxed:.2f}''')
        self.lbl_kpi_cost.config(text = f'''综合到岸总成本: ¥{total_cost:.2f}''')

    
    def on_add_item_row(self):
        self.open_item_edit_dialog(None)

    
    def on_edit_item_row(self):
        sel = self.tree_items.selection()
        if not sel:
            messagebox.showwarning('提示', '请先选中要编辑的货物明细行！')
            return None
        self.open_item_edit_dialog(sel[0])

    
    def on_delete_item_row(self):
        sel = self.tree_items.selection()
        if not sel:
            return None
        for s in sel:
            self.tree_items.delete(s)
        self.recalculate_single_tab()

    
    def on_clear_items(self):
        if messagebox.askyesno('确认', '确定要清空当前所有明细吗？'):
            for s in self.tree_items.get_children():
                self.tree_items.delete(s)
            self.recalculate_single_tab()
            return None

    
    def open_item_edit_dialog(self, item_id):
        dlg = tk.Toplevel(self.root)
        dlg.title('编辑型号明细' if item_id else '新增型号明细')
        dlg.geometry('440x480')
        dlg.resizable(False, False)
        dlg.transient(self.root)
        dlg.grab_set()
        initial_vals = self.tree_items.item(item_id, 'values') if item_id else [
            '电焊机',
            '',
            '',
            '1',
            '1',
            '1',
            '0.00',
            '0.00',
            '0.00',
            '0.00',
            '0.00',
            '0.00']
        f = tk.Frame(dlg, padx = 20, pady = 15)
        f.pack(fill = 'both', expand = True)
        fields = [
            ('品名:', initial_vals[0]),
            ('原始型号/品目名:', initial_vals[2] or initial_vals[1]),
            ('箱数:', initial_vals[3]),
            ('装数 (台/箱):', initial_vals[4]),
            ('出货数量 (台):', initial_vals[5]),
            ('含税单价 (¥):', initial_vals[6]),
            ('体积 (CBM):', initial_vals[9])]
        entries = []
        for i, (lbl, val) in enumerate(fields):
            tk.Label(f, text = lbl, font = ('微软雅黑', 9, 'bold' if i in (1, 5, 6) else 'normal')).grid(row = i, column = 0, sticky = 'w', pady = 5)
            ent = ttk.Entry(f, width = 24)
            ent.insert(0, str(val))
            ent.grid(row = i, column = 1, sticky = 'ew', pady = 5, padx = 6)
            entries.append(ent)
        f.columnconfigure(1, weight = 1)
        
        def save():
            cat = entries[0].get().strip()
            raw_model = entries[1].get().strip()
            boxes = core.safe_int(entries[2].get(), 1)
            pack_qty = core.safe_int(entries[3].get(), 1)
            qty = core.safe_int(entries[4].get(), boxes * pack_qty)
            price = core.safe_float(entries[5].get(), 0.0)
            cbm = core.safe_float(entries[6].get(), 0.0)
            if not raw_model:
                messagebox.showerror('错误', '型号规格不能为空！', parent = dlg)
                return None
            clean_code = core.sanitize_item_code(raw_model)
            total_p = round(price * qty, 2)
            untaxed_u = round(price / 1.13, 4)
            untaxed_tot = round(untaxed_u * qty, 2)
            vals = (cat or raw_model, clean_code, raw_model, boxes, pack_qty, qty, f'''{price:.2f}''', f'''{total_p:.2f}''', f'''{untaxed_tot:.2f}''', f'''{cbm:.4f}''', '0.00', '0.00')
            if item_id:
                self.tree_items.item(item_id, values = vals)
            else:
                self.tree_items.insert('', 'end', values = vals)
            self.recalculate_single_tab()
            dlg.destroy()

        btn_box = tk.Frame(f)
        btn_box.grid(row = len(fields), column = 0, columnspan = 2, pady = 16)
        ttk.Button(btn_box, text = '确定保存', style = 'Primary.TButton', command = save).pack(side = 'left', padx = 8)
        ttk.Button(btn_box, text = '取消', command = dlg.destroy).pack(side = 'left', padx = 8)

    
    def on_export_single_container(self):
        c_info = {
            'trade_co': self.var_s_company.get().strip(),
            'contract_no': self.var_s_contract.get().strip(),
            'container_no': self.var_s_container.get().strip(),
            'date': self.var_s_date.get().strip(),
            'seq_no': core.safe_int(self.var_s_seq.get(), 1),
            'rate': core.safe_float(self.var_s_rate.get(), 6.7322),
            'warehouse': self.var_s_warehouse.get().strip() or '100',
            'ocean_usd': core.safe_float(self.var_s_ocean.get(), 0.0),
            'truck_rmb': core.safe_float(self.var_s_truck.get(), 0.0),
            'insurance_usd': core.safe_float(self.var_s_insurance.get(), 0.0),
            'agency_rmb': core.safe_float(self.var_s_agency.get(), 0.0) }
        items = []
        for item_id in self.tree_items.get_children():
            vals = self.tree_items.item(item_id, 'values')
            items.append({
                'cat': vals[0],
                'model': vals[1],
                'raw_model': vals[2],
                'boxes': core.safe_int(vals[3]),
                'pack_qty': core.safe_int(vals[4]),
                'qty': core.safe_int(vals[5]),
                'unit_price': core.safe_float(vals[6]),
                'total_price': core.safe_float(vals[7]),
                'untaxed_total': core.safe_float(vals[8]),
                'cbm': core.safe_float(vals[9]) })
        if not items:
            messagebox.showerror('错误', '当前货柜明细为空，无法生成上传表！')
            return None
        out_dir = self.var_output_dir.get().strip()
        os.makedirs(out_dir, exist_ok = True)
        filename = f'''{c_info['contract_no']}_{c_info['trade_co']}_ECOUNT全套上传表.xlsx'''
        filename = re.sub('[\\\\/*?:"<>|]', '_', filename)
        out_file = os.path.join(out_dir, filename)
        
        try:
            core.build_erp_workbook(c_info, items, existing_master = self.master_items, output_path = out_file, template_paths = self.template_config, export_individual_files = self.var_export_individual.get(), log_func = self.log)
            msg_succ = f'''成功生成 ECOUNT ERP 全套上传表：\n{out_file}'''
            if self.var_export_individual.get():
                msg_succ += '\n\n(已同步生成对应的独立上传单表文件)'
            res = messagebox.askyesno('导出成功', f'''{msg_succ}\n\n是否立即打开该文件？''')
            if res:
                if sys.platform == 'win32':
                    os.startfile(out_file)
                    return None
                subprocess.run([
                    'xdg-open',
                    out_file])
                return None
            return None
        except Exception as e:
            self.log(f'''[Error] 导出失败: {e}''', 'err')
            messagebox.showerror('导出失败', str(e))
            return None


    
    def on_batch_generate(self):
        selected_indices = []
        for item_id in self.tree_containers.get_children():
            vals = self.tree_containers.item(item_id, 'values')
            if vals[0] == '☑ 是':
                selected_indices.append(int(item_id))
        if not selected_indices:
            messagebox.showwarning('提示', '请至少勾选一个要生成的货柜！')
            return None
        source_file = self.var_source_file.get().strip()
        out_dir = self.var_output_dir.get().strip()
        threading.Thread(target = self._async_batch_generate, args = (source_file, out_dir, selected_indices), daemon = True).start()

    
    def _async_batch_generate(self, source_file, out_dir, selected_indices):
        self.root.after(0, (lambda : self.set_busy_state(True, '正在批量生成上传工作簿...')))
        
        try:
            
            def prog_cb(cur, total):
                pct = int((cur / total) * 100)
                self.root.after(0, (lambda : self.progress_bar.config(value = pct)))
                self.root.after(0, (lambda : self.lbl_progress.config(text = f'''正在生成 {cur}/{total} ({pct}%)''')))

            files = core.batch_process_containers(source_file, output_dir = out_dir, selected_indices = selected_indices, master_data = self.master_items, custom_mapping = self.header_mapping, template_paths = self.template_config, export_individual_files = self.var_export_individual.get(), progress_callback = prog_cb, log_callback = (lambda m: self.root.after(0, self.log, m)))
            self.last_generated_files = files
            
            def on_done():
                self.set_busy_state(False, f'''已完成生成 {len(files)} 个工作簿''')
                res = messagebox.askyesno('生成完成', f'''🎉 批量生成完成！共生成 {len(files)} 份 ECOUNT 标准上传工作簿。\n\n保存在: {out_dir}\n\n是否立即打开输出目录？''')
                if res:
                    self.on_open_output_dir()
                    return None

            self.root.after(0, on_done)
            return None
        except Exception as e:
            err_msg = str(e)
            self.root.after(0, (lambda msg=err_msg: self.log(f'''[Error] 批量生成失败: {msg}''', 'err')))
            self.root.after(0, (lambda msg=err_msg: messagebox.showerror('批量生成失败', msg)))
            self.root.after(0, (lambda : self.set_busy_state(False, '生成出错')))
            return None


    
    def set_busy_state(self, is_busy, status_msg = ''):
        if is_busy:
            self.progress_bar.config(mode = 'indeterminate')
            self.progress_bar.start(10)
            self.btn_batch_generate.config(state = 'disabled')
        else:
            self.progress_bar.stop()
            self.progress_bar.config(mode = 'determinate', value = 100)
            self.btn_batch_generate.config(state = 'normal')
        self.lbl_progress.config(text = status_msg)

    
    def refresh_master_table(self, filter_kw = ''):
        for item in self.tree_master.get_children():
            self.tree_master.delete(item)
        kw = filter_kw.strip().upper()
        count = 0
        for r in self.master_items:
            code = str(r[0]).strip()
            name = str(r[1]).strip()
            if kw and kw not in code.upper() and kw not in name.upper():
                continue
            self.tree_master.insert('', 'end', values = r)
            count += 1
        self.lbl_master_count.config(text = f'''品目总数: {len(self.master_items)} (显示 {count})''')

    
    def on_filter_master_items(self, *args):
        self.refresh_master_table(self.var_master_search.get())

    
    def on_load_external_master(self):
        path = filedialog.askopenfilename(title = '选择 ESA009M 或已有上传工作簿', filetypes = [
            ('Excel 文件', '*.xlsx *.xls'),
            ('所有文件', '*.*')])
        if path:
            
            try:
                new_master = core.load_master_from_excel(path)
                self.master_items = new_master
                self.refresh_master_table()
                self.log(f'''[OK] 成功从 {os.path.basename(path)} 载入 {len(new_master)} 条品目编码主数据！''', 'ok')
                messagebox.showinfo('成功', f'''成功载入 {len(new_master)} 条品目主数据！''')
                return None
            except Exception as e:
                self.log(f'''[Error] 载入主数据失败: {e}''', 'err')
                messagebox.showerror('载入失败', str(e))
                return None


    
    def on_add_master_item(self):
        dlg = tk.Toplevel(self.root)
        dlg.title('新增品目主数据')
        dlg.geometry('400x380')
        dlg.resizable(False, False)
        dlg.transient(self.root)
        dlg.grab_set()
        f = tk.Frame(dlg, padx = 20, pady = 15)
        f.pack(fill = 'both', expand = True)
        fields = [
            '型号/名称 (*):',
            '品目类型:',
            '规格信息:',
            '装数:',
            '体积 (CBM):',
            '重量 (kg):']
        defaults = [
            '',
            '[商品]',
            '',
            '1',
            '',
            '']
        entries = []
        for i, (lbl, d_val) in enumerate(zip(fields, defaults)):
            tk.Label(f, text = lbl, font = ('微软雅黑', 9, 'bold' if '(*)' in lbl else 'normal')).grid(row = i, column = 0, sticky = 'w', pady = 5)
            ent = ttk.Entry(f, width = 22)
            ent.insert(0, d_val)
            ent.grid(row = i, column = 1, sticky = 'ew', pady = 5, padx = 6)
            entries.append(ent)
        f.columnconfigure(1, weight = 1)
        
        def save():
            raw_code = entries[0].get().strip()
            t = entries[1].get().strip() or '[商品]'
            spec = entries[2].get().strip() or raw_code
            pack_qty = entries[3].get().strip()
            cbm = entries[4].get().strip()
            weight = entries[5].get().strip()
            if not raw_code:
                messagebox.showerror('错误', '品目编码/型号不能为空！', parent = dlg)
                return None
            clean_code = core.sanitize_item_code(raw_code)
            new_row = (clean_code, raw_code, t, spec, '', '', 'YES', 1, pack_qty, cbm, weight)
            self.master_items.append(new_row)
            self.refresh_master_table()
            self.log(f'''[OK] 新增品目编码主数据: {clean_code} (原名: {raw_code})''', 'ok')
            dlg.destroy()

        btn_box = tk.Frame(f)
        btn_box.grid(row = len(fields), column = 0, columnspan = 2, pady = 16)
        ttk.Button(btn_box, text = '保存', style = 'Primary.TButton', command = save).pack(side = 'left', padx = 8)
        ttk.Button(btn_box, text = '取消', command = dlg.destroy).pack(side = 'left', padx = 8)

    
    def on_delete_master_item(self):
        sel = self.tree_master.selection()
        if not sel:
            return None
        if messagebox.askyesno('确认', f'''确定删除选中的 {len(sel)} 个品目主数据吗？'''):
            for s in sel:
                vals = self.tree_master.item(s, 'values')
                self.master_items = [ x for x in self.master_items if str(x[0]) != str(vals[0]) ]
            self.refresh_master_table()
            return None

    
    def refresh_counterparty_table(self):
        for item in self.tree_cp.get_children():
            self.tree_cp.delete(item)
        for key, val in self.counterparty_map.items():
            vals = (key, val.get('rmb_code', ''), val.get('rmb_name', ''), val.get('usd_code', ''), val.get('usd_name', ''))
            self.tree_cp.insert('', 'end', values = vals)
        self.combo_s_company['values'] = list(self.counterparty_map.keys())

    
    def on_add_counterparty(self):
        self.open_counterparty_dialog(None)

    
    def on_edit_counterparty(self):
        sel = self.tree_cp.selection()
        if not sel:
            messagebox.showwarning('提示', '请选中要编辑的往来单位！')
            return None
        vals = self.tree_cp.item(sel[0], 'values')
        self.open_counterparty_dialog(vals[0])

    
    def on_delete_counterparty(self):
        sel = self.tree_cp.selection()
        if not sel:
            return None
        vals = self.tree_cp.item(sel[0], 'values')
        key = vals[0]
        if messagebox.askyesno('确认', f'''确定删除往来单位 [{key}] 的配置吗？'''):
            if key in self.counterparty_map:
                del self.counterparty_map[key]
                core.save_counterparty_map(self.counterparty_map)
                self.refresh_counterparty_table()
                self.log(f'''[*] 已删除往来单位配置: {key}''')
                return None
        return None

    
    def on_save_counterparty_config(self):
        if core.save_counterparty_map(self.counterparty_map):
            messagebox.showinfo('成功', '往来单位映射配置已保存到本地！')
            self.log('[OK] 往来单位映射配置已成功写入本地 JSON 文件', 'ok')
            return None

    
    def open_counterparty_dialog(self, edit_key):
        dlg = tk.Toplevel(self.root)
        dlg.title('编辑往来单位' if edit_key else '新增往来单位')
        dlg.geometry('450x360')
        dlg.resizable(False, False)
        dlg.transient(self.root)
        dlg.grab_set()
        f = tk.Frame(dlg, padx = 20, pady = 15)
        f.pack(fill = 'both', expand = True)
        info = self.counterparty_map.get(edit_key, { }) if edit_key else { }
        fields = [
            ('公司简称/关键词 (*):', edit_key or ''),
            ('人民币往来单位编码 (*):', info.get('rmb_code', '')),
            ('人民币往来单位全称 (*):', info.get('rmb_name', '')),
            ('美金往来单位编码 (*):', info.get('usd_code', '')),
            ('美金往来单位全称 (*):', info.get('usd_name', ''))]
        entries = []
        for i, (lbl, val) in enumerate(fields):
            tk.Label(f, text = lbl, font = ('微软雅黑', 9, 'bold' if '(*)' in lbl else 'normal')).grid(row = i, column = 0, sticky = 'w', pady = 5)
            ent = ttk.Entry(f, width = 26)
            ent.insert(0, val)
            ent.grid(row = i, column = 1, sticky = 'ew', pady = 5, padx = 6)
            entries.append(ent)
        f.columnconfigure(1, weight = 1)
        
        def save():
            key = entries[0].get().strip()
            rmb_c = entries[1].get().strip()
            rmb_n = entries[2].get().strip()
            usd_c = entries[3].get().strip()
            usd_n = entries[4].get().strip()
            if not (key and rmb_c) or not usd_c:
                messagebox.showerror('错误', '公司简称及往来单位编码不能为空！', parent = dlg)
                return None
            self.counterparty_map[key] = {
                'rmb_code': rmb_c,
                'rmb_name': rmb_n or f'''浙江{key}有限公司''',
                'usd_code': usd_c,
                'usd_name': usd_n or f'''浙江{key}有限公司USD''' }
            core.save_counterparty_map(self.counterparty_map)
            self.refresh_counterparty_table()
            self.log(f'''[OK] 往来单位配置更新: {key} -> RMB[{rmb_c}] USD[{usd_c}]''', 'ok')
            dlg.destroy()

        btn_box = tk.Frame(f)
        btn_box.grid(row = len(fields), column = 0, columnspan = 2, pady = 16)
        ttk.Button(btn_box, text = '保存', style = 'Primary.TButton', command = save).pack(side = 'left', padx = 8)
        ttk.Button(btn_box, text = '取消', command = dlg.destroy).pack(side = 'left', padx = 8)



def main():
    root = tk.Tk()
    app = ContainerProcessingApp(root)
    root.mainloop()

if __name__ == '__main__':
    main()
